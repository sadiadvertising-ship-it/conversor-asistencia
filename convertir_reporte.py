"""
Conversor del "Reporte de Eventos" de asistencia.

Separa la columna "Fecha/Hora" (texto MM/DD/YYYY HH:MM) en dos columnas:
"Fecha" y "Hora", como valores reales de Excel. El archivo original NUNCA
se modifica: se genera <nombre>_convertido.xlsx en la misma carpeta.

Uso:
  - Arrastrar uno o varios reportes .xlsx sobre "CONVERTIR (arrastre aqui el reporte).bat"
  - o doble clic en el .bat y elegir el archivo en la ventana
  - o: python convertir_reporte.py reporte1.xlsx [reporte2.xlsx ...]
"""

import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COL_FECHA_HORA = "Fecha/Hora"
COLUMNAS_ORIGEN = [
    "ID de Usuario",
    "Nombre Completo",
    COL_FECHA_HORA,
    "Código de Trabajo",
    "Tipo de Registro",
]
FORMATOS = ["%m/%d/%Y %H:%M", "%m/%d/%Y %H:%M:%S"]
FILAS_BUSQUEDA_ENCABEZADO = 60


class ErrorReporte(Exception):
    """Error con un mensaje entendible para el usuario."""


def texto(valor):
    return "" if valor is None else str(valor).strip()


def buscar_encabezado(filas):
    """Devuelve (indice_fila, {nombre_columna: indice_columna})."""
    for i, fila in enumerate(filas[:FILAS_BUSQUEDA_ENCABEZADO]):
        nombres = [texto(c) for c in fila]
        if COL_FECHA_HORA in nombres:
            mapa = {n: nombres.index(n) for n in COLUMNAS_ORIGEN if n in nombres}
            return i, mapa
    raise ErrorReporte(
        f"No se encontró la columna '{COL_FECHA_HORA}' en las primeras "
        f"{FILAS_BUSQUEDA_ENCABEZADO} filas. ¿Es el reporte correcto?"
    )


def buscar_rango(filas, hasta):
    for fila in filas[:hasta]:
        for c in fila:
            if texto(c).startswith("Desde"):
                return texto(c)
    return ""


def separar(valor):
    """Devuelve (fecha, hora, con_segundos) o None si no cumple el formato."""
    if isinstance(valor, datetime):
        return valor.date(), valor.time(), valor.second != 0
    t = texto(valor)
    for fmt in FORMATOS:
        try:
            dt = datetime.strptime(t, fmt)
            return dt.date(), dt.time(), fmt.endswith("%S")
        except ValueError:
            continue
    return None


def leer_reporte(ruta):
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        filas = [list(f) for f in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    i_enc, mapa = buscar_encabezado(filas)
    rango = buscar_rango(filas, i_enc)

    def celda(fila, nombre):
        j = mapa.get(nombre)
        return fila[j] if j is not None and j < len(fila) else None

    registros, revisar = [], []
    for n, fila in enumerate(filas[i_enc + 1:], start=i_enc + 2):
        fh = celda(fila, COL_FECHA_HORA)
        nombre = celda(fila, "Nombre Completo")
        if texto(fh) == COL_FECHA_HORA:
            continue  # encabezado repetido
        if texto(fh) == "":
            if texto(nombre):
                revisar.append((n, "Fila con nombre pero sin Fecha/Hora", fila))
            continue
        reg = {k: celda(fila, k) for k in COLUMNAS_ORIGEN}
        reg["fila_origen"] = n
        reg["partes"] = separar(fh)
        if reg["partes"] is None:
            revisar.append((n, f"Fecha/Hora con formato no reconocido: {fh!r}", fila))
        registros.append(reg)
    return registros, revisar, rango


def ajustar_anchos(ws, minimo=10, maximo=45):
    for col in ws.columns:
        largo = max(len(texto(c.value)) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(minimo, min(maximo, largo + 2))


def estilo_encabezado(ws):
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        c.alignment = Alignment(horizontal="center")


def escribir_salida(ruta_origen, registros, revisar, rango, ruta_salida):
    wb = Workbook()
    ws = wb.active
    ws.title = "Registros"
    ws.append([
        "ID de Usuario", "Nombre Completo", "Fecha", "Hora",
        "Fecha/Hora original", "Código de Trabajo", "Tipo de Registro",
    ])
    amarillo = PatternFill("solid", fgColor="FFF2CC")
    for reg in registros:
        partes = reg["partes"]
        fecha, hora = (partes[0], partes[1]) if partes else (None, None)
        ws.append([
            reg["ID de Usuario"], reg["Nombre Completo"], fecha, hora,
            reg[COL_FECHA_HORA], reg["Código de Trabajo"], reg["Tipo de Registro"],
        ])
        r = ws.max_row
        ws.cell(r, 3).number_format = "DD/MM/YYYY"
        ws.cell(r, 4).number_format = "HH:MM:SS" if partes and partes[2] else "HH:MM"
        if partes is None:
            for c in ws[r]:
                c.fill = amarillo
    estilo_encabezado(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ajustar_anchos(ws)

    if revisar:
        wr = wb.create_sheet("Revisar")
        wr.append(["Fila en el original", "Motivo", "Contenido de la fila"])
        for n, motivo, fila in revisar:
            wr.append([n, motivo, " | ".join(texto(c) for c in fila if texto(c))])
        estilo_encabezado(wr)
        ajustar_anchos(wr, maximo=80)

    convertidos = sum(1 for r in registros if r["partes"])
    rs = wb.create_sheet("Resumen")
    for fila in (
        ("Archivo origen", str(ruta_origen)),
        ("Rango del reporte", rango),
        ("Convertido el", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Registros leídos", len(registros)),
        ("Registros escritos", ws.max_row - 1),
        ("Fecha/Hora separadas correctamente", convertidos),
        ("Filas a revisar", len(revisar)),
    ):
        rs.append(fila)
    for c in rs["A"]:
        c.font = Font(bold=True)
    ajustar_anchos(rs, maximo=90)

    wb.save(ruta_salida)
    return convertidos


def procesar(ruta):
    """Convierte un reporte y devuelve un diccionario con el resultado.

    Lanza ErrorReporte con un mensaje para el usuario si algo falla.
    """
    ruta = Path(ruta).resolve()
    if ruta.suffix.lower() != ".xlsx":
        raise ErrorReporte("El archivo debe ser .xlsx")
    if ruta.stem.endswith("_convertido"):
        raise ErrorReporte("Este archivo ya es un resultado convertido.")

    registros, revisar, rango = leer_reporte(ruta)
    ruta_salida = ruta.with_name(f"{ruta.stem}_convertido.xlsx")
    try:
        convertidos = escribir_salida(ruta, registros, revisar, rango, ruta_salida)
    except PermissionError:
        raise ErrorReporte(
            f"No se pudo guardar {ruta_salida.name}. Probablemente está abierto "
            "en Excel: ciérrelo e intente de nuevo."
        ) from None
    return {
        "origen": ruta,
        "salida": ruta_salida,
        "rango": rango,
        "registros": registros,
        "convertidos": convertidos,
        "revisar": revisar,
    }


def convertir(ruta):
    print(f"\nProcesando: {Path(ruta).name}")
    try:
        r = procesar(ruta)
    except ErrorReporte as e:
        print(f"  ERROR: {e}")
        return False
    print(f"  {r['rango']}")
    print(f"  Registros: {len(r['registros'])}  |  Separados OK: {r['convertidos']}"
          f"  |  A revisar: {len(r['revisar'])}")
    print(f"  Guardado en: {r['salida']}")
    if r["revisar"]:
        print("  ATENCION: revise la hoja 'Revisar' del archivo convertido.")
    return not r["revisar"]


def pedir_archivos():
    try:
        import tkinter as tk
        from tkinter import filedialog
    except ImportError:
        return []
    raiz = tk.Tk()
    raiz.withdraw()
    archivos = filedialog.askopenfilenames(
        title="Seleccione el/los reporte(s) de eventos",
        filetypes=[("Excel", "*.xlsx")],
        initialdir=str(Path.home() / "Downloads"),
    )
    raiz.destroy()
    return list(archivos)


def main():
    archivos = sys.argv[1:] or pedir_archivos()
    if not archivos:
        print("No se seleccionó ningún archivo.")
        return 1
    todo_ok = True
    for a in archivos:
        try:
            todo_ok &= convertir(a)
        except Exception as e:
            print(f"  ERROR: {e}")
            todo_ok = False
    return 0 if todo_ok else 1


if __name__ == "__main__":
    sys.exit(main())
